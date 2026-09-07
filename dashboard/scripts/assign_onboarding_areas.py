#!/usr/bin/env python3
"""
assign_onboarding_areas.py
Assigns area (ARB / Non-ARB / Minor) to all Onboarding features in features-master.xlsx.
"""
import os, shutil
from datetime import date
from openpyxl import load_workbook

# Area assignments: feature_id → area
AREA_MAP = {
    # ARB — Standard Al Rajhi Bank customer onboarding
    "ONB-001": "ARB",   # Digital Account Opening (Saudi)
    "ONB-003": "ARB",   # Update National Address
    "ONB-004": "ARB",   # Session Upgrade & OTP
    "ONB-007": "ARB",   # Expired ID Handling
    "ONB-010": "ARB",   # FATCA/CRS KYC Update
    "ONB-014": "ARB",   # Splash Screen & App Launch
    "ONB-015": "ARB",   # Login Screen
    "ONB-016": "ARB",   # Add/Create Password
    "ONB-017": "ARB",   # Face ID / Biometric Setup
    "ONB-018": "ARB",   # Change Password
    "ONB-020": "ARB",   # KYC Confirmation Screens
    "ONB-022": "ARB",   # Progress Indicator / Stepper
    "ONB-023": "ARB",   # Pre-filled KYC from National ID
    "ONB-024": "ARB",   # Onboarding Checklist / Guide
    "ONB-025": "ARB",   # First Engagement
    "ONB-026": "ARB",   # Fast Onboarding
    "ONB-039": "ARB",   # Guest Mode
    "ONB-040": "ARB",   # Tutorials / How-To Guides
    "ONB-041": "ARB",   # How to Deposit Guide
    "ONB-042": "ARB",   # How to Trade Guide
    "ONB-043": "ARB",   # Data Consent Management
    "ONB-044": "ARB",   # T&C Version Tracking
    "ONB-045": "ARB",   # Passwordless Authentication
    "SAU-005": "ARB",   # Block Codes (AML)

    # Non-ARB — GCC, foreign, corporate, IBKR onboarding
    "ONB-005": "Non-ARB",  # GCC Customer Onboarding
    "ONB-006": "Non-ARB",  # Family Onboarding (Non-Saudi)
    "ONB-008": "Non-ARB",  # GCC Customer ID Update
    "ONB-009": "Non-ARB",  # Foreign Customer ID Verification
    "ONB-013": "Non-ARB",  # IBKR Intl Brokerage Onboarding
    "ONB-019": "Non-ARB",  # Non-ARB Bank Account Entry
    "ONB-021": "Non-ARB",  # Family Members Overview
    "ONB-028": "Non-ARB",  # Corporate — Other Local Entities
    "ONB-029": "Non-ARB",  # Corporate — Approval Flow
    "ONB-030": "Non-ARB",  # Corporate — Under Review
    "ONB-031": "Non-ARB",  # Corporate — Additional Requirements
    "ONB-032": "Non-ARB",  # Corporate — Rejected
    "ONB-033": "Non-ARB",  # Corporate — Local Company Onboarding
    "ONB-034": "Non-ARB",  # Corporate — Profile & Settings
    "ONB-035": "Non-ARB",  # Corporate — Requirements Upload
    "ONB-036": "Non-ARB",  # Corporate — Terms & Conditions
    "ONB-037": "Non-ARB",  # Corporate — Status Tracking
    "ONB-038": "Non-ARB",  # Corporate — Pending Actions

    # Minor — Minor account onboarding and guardian controls
    "ONB-002": "Minor",  # Minor Account Onboarding
    "ONB-011": "Minor",  # Minor Trading Restrictions
    "ONB-012": "Minor",  # Guardian Trading for Minor
    "ONB-027": "Minor",  # Minor Account + Guardian Controls
}

here = os.path.dirname(os.path.abspath(__file__))
dashboard_dir = os.path.join(here, "..")
xlsx_path = os.path.join(dashboard_dir, "features-master.xlsx")
backup_path = xlsx_path + f".bak-{date.today().isoformat()}-onb-areas"

shutil.copy2(xlsx_path, backup_path)
print(f"Backup: {backup_path}")

wb = load_workbook(xlsx_path)
ws = wb["Features"]

updated = 0
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if not fid:
        continue
    fid = str(fid).strip()
    if fid in AREA_MAP:
        ws.cell(row=r, column=10, value=AREA_MAP[fid])  # col J = area
        updated += 1

wb.save(xlsx_path)
wb.close()
print(f"Updated {updated} / {len(AREA_MAP)} Onboarding features with area assignments")
print(f"  ARB:     {sum(1 for v in AREA_MAP.values() if v == 'ARB')}")
print(f"  Non-ARB: {sum(1 for v in AREA_MAP.values() if v == 'Non-ARB')}")
print(f"  Minor:   {sum(1 for v in AREA_MAP.values() if v == 'Minor')}")
print(f"Saved: {xlsx_path}")
