#!/usr/bin/env python3
"""
duplicate_shared_onboarding.py
Duplicates shared onboarding steps across all user types (Local, Global, Corporate, Minor).
Each area gets the complete flow — shared steps appear under every area.
"""
import os, shutil, json
from datetime import date
from openpyxl import load_workbook

# Shared features: these exist in ARB but apply to ALL user types
SHARED_IDS = [
    "ONB-004",   # Session Upgrade & OTP
    "ONB-010",   # FATCA/CRS KYC Update
    "ONB-014",   # Splash Screen & App Launch
    "ONB-015",   # Login Screen
    "ONB-016",   # Add/Create Password
    "ONB-017",   # Face ID / Biometric Setup
    "ONB-018",   # Change Password
    "ONB-020",   # KYC Confirmation Screens
    "ONB-022",   # Progress Indicator / Stepper
    "ONB-024",   # Onboarding Checklist / Guide
    "ONB-025",   # First Engagement
    "ONB-039",   # Guest Mode
    "ONB-040",   # Tutorials / How-To Guides
    "ONB-041",   # How to Deposit Guide
    "ONB-042",   # How to Trade Guide
    "ONB-043",   # Data Consent Management
    "ONB-044",   # T&C Version Tracking
    "ONB-045",   # Passwordless Authentication
]

# Target areas to duplicate into (already exists in ARB)
TARGET_AREAS = ["Local", "Global", "Corporate", "Minor"]

here = os.path.dirname(os.path.abspath(__file__))
dashboard_dir = os.path.join(here, "..")
xlsx_path = os.path.join(dashboard_dir, "features-master.xlsx")
backup_path = xlsx_path + f".bak-{date.today().isoformat()}-dup-shared"

shutil.copy2(xlsx_path, backup_path)
print(f"Backup: {backup_path}")

wb = load_workbook(xlsx_path)
ws = wb["Features"]

# Read shared features data from xlsx
shared_data = {}
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if not fid:
        continue
    fid = str(fid).strip()
    if fid in SHARED_IDS:
        shared_data[fid] = {
            "name": ws.cell(row=r, column=2).value,
            "status": ws.cell(row=r, column=4).value,
            "definition": ws.cell(row=r, column=5).value,
            "brd": ws.cell(row=r, column=6).value,
            "competitors": ws.cell(row=r, column=7).value,
            "priority": ws.cell(row=r, column=8).value,
            "tags": ws.cell(row=r, column=9).value,
        }

print(f"Found {len(shared_data)} shared features to duplicate across {len(TARGET_AREAS)} areas")

# Find last data row
last_row = 4
for r in range(ws.max_row, 3, -1):
    if any(ws.cell(row=r, column=c).value for c in range(1, 5)):
        last_row = r
        break

# Get next ONB ID
onb_ids = []
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if fid and str(fid).strip().startswith("ONB-"):
        try:
            onb_ids.append(int(str(fid).strip().split("-")[1]))
        except ValueError:
            pass
next_id = max(onb_ids) + 1

next_row = last_row + 1
total_added = 0

for area in TARGET_AREAS:
    area_count = 0
    for orig_id in SHARED_IDS:
        if orig_id not in shared_data:
            continue
        sd = shared_data[orig_id]
        fid = f"ONB-{next_id:03d}"
        ws.cell(row=next_row, column=1, value=fid)
        ws.cell(row=next_row, column=2, value=sd["name"])
        ws.cell(row=next_row, column=3, value="Onboarding")
        ws.cell(row=next_row, column=4, value=sd["status"])
        if sd["definition"]:
            ws.cell(row=next_row, column=5, value=sd["definition"])
        if sd["brd"]:
            ws.cell(row=next_row, column=6, value=sd["brd"])
        if sd["competitors"]:
            ws.cell(row=next_row, column=7, value=sd["competitors"])
        if sd["priority"] is not None:
            ws.cell(row=next_row, column=8, value=sd["priority"])
        if sd["tags"]:
            ws.cell(row=next_row, column=9, value=sd["tags"])
        ws.cell(row=next_row, column=10, value=area)

        next_id += 1
        next_row += 1
        area_count += 1
        total_added += 1
    print(f"  {area}: +{area_count} features (ONB-{next_id - area_count:03d} to ONB-{next_id - 1:03d})")

wb.save(xlsx_path)
wb.close()
print(f"\nTotal added: {total_added}")
print(f"Saved: {xlsx_path}")
