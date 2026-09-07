#!/usr/bin/env python3
"""
assign_platform_areas.py
Files the five features that carried no area, and widens Education to Education/Social.

Ahmed's ruling, 2026-09-06, after the Features Map deck surfaced 5 of 865 rows
rendering as "Unassigned" on the Platform & Managed Products slide:

    SAU-006  TILA                          -> Platform
    XJ-049   ARC Chatbot                   -> Platform
    XJ-068   Cross-Product Gain-Loss Chart -> Platform
    XJ-069   Global Last Transaction View  -> Cash Management
    SAU-099  Social / Copy Trading         -> Education/Social

The Education area is renamed Education/Social so it can carry SAU-099 alongside
the three existing education rows. Column J (10) is the area column, data starts
at row 4 — same contract as assign_onboarding_areas.py.
"""
import os
import shutil
from datetime import date

from openpyxl import load_workbook

AREA_MAP = {
    "SAU-006": "Platform",           # TILA
    "XJ-049":  "Platform",           # ARC Chatbot
    "XJ-068":  "Platform",           # Cross-Product Gain-Loss Chart
    "XJ-069":  "Cash Management",    # Global Last Transaction View
    "SAU-099": "Education/Social",   # Social / Copy Trading
}
RENAME = {"Education": "Education/Social"}

here = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(here, "..", "features-master.xlsx")
backup = xlsx + f".bak-{date.today().isoformat()}-platform-areas"
shutil.copy2(xlsx, backup)
print(f"Backup: {os.path.basename(backup)}")

wb = load_workbook(xlsx)
ws = wb["Features"]

assigned, renamed = 0, 0
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if not fid:
        continue
    fid = str(fid).strip()
    cur = ws.cell(row=r, column=10).value
    if fid in AREA_MAP:
        ws.cell(row=r, column=10, value=AREA_MAP[fid])
        print(f"  {fid:9s} area {cur!r} -> {AREA_MAP[fid]!r}")
        assigned += 1
    elif cur and str(cur).strip() in RENAME:
        new = RENAME[str(cur).strip()]
        ws.cell(row=r, column=10, value=new)
        renamed += 1

wb.save(xlsx)
wb.close()
print(f"\nAssigned {assigned}/{len(AREA_MAP)} areas · renamed {renamed} Education rows")
