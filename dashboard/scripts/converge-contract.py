#!/usr/bin/env python3
"""Converge the Saudi Market area to baseline v1.4. Atomic, keyed by ID."""
import shutil
from datetime import date
from pathlib import Path
import openpyxl

ROOT = Path('/Users/ahmedalghamdi/Claude/ARC')
X = ROOT / 'projects/features-map/dashboard/features-master.xlsx'
shutil.copyfile(X, X.with_name(X.name + '.bak-2026-08-17-market-v1.4'))
shutil.copyfile(X, '/private/tmp/claude-501/-Users-ahmedalghamdi-Claude-ARC/64e640d3-d9cf-4e26-a0ac-677bd5d66a89/scratchpad/pre_v14.xlsx')
print('backup -> features-master.xlsx.bak-2026-08-17-market-v1.4')

MP, SP = 'Market Page', 'Stock Page'
C = [
 (MP,'-','Market Chart',None),(MP,'-','Market Hours',None),(MP,'-','Market Summary',None),
 (MP,'-','Stock Search','SAU-123'),
 (MP,'-','Market Sectors','SAU-095'),(MP,'Market Sectors','Sector Details',None),
 (MP,'Market Sectors','Volatility','SAU-025'),
 (MP,'-','Stocks',None),(MP,'Stocks','Trending','SAU-018'),(MP,'Stocks','Top Gainers',None),
 (MP,'Stocks','Top Losers',None),(MP,'Stocks','Most Active by Quantity',None),
 (MP,'Stocks','Most Active by Value',None),(MP,'Stocks','Top Traders',None),
 (MP,'Stocks','Sharia Stock','SAU-124'),
 (MP,'-','Economic Calendar',None),(MP,'-','Earnings Calendar','SAU-107'),
 (MP,'-','Dividend Calendar','SAU-116'),(MP,'-','News','SAU-067'),
 (MP,'News','News Screener','XJ-006'),(MP,'-','Stock Screener','SAU-039'),
 (MP,'Stock Screener','Saved Themes','SAU-144'),(MP,'-','Analysis Changes','SAU-076'),
 (MP,'-','ETF View','SAU-110'),(MP,'-','Whale/Institutional Trades','SAU-111'),
 (MP,'-','Unusual Activity','SAU-112'),
 (SP,'-','Why Is It Moving','SAU-064'),(SP,'-','Earnings Calendar',None),(SP,'-','Trending Tag',None),
 (SP,'-','Stock Details',None),(SP,'Stock Details','Market Depth','SAU-038'),
 (SP,'Stock Details','Traders Summary',None),(SP,'Stock Details','Key Statistics',None),
 (SP,'-','Stock News','SAU-132'),(SP,'-','Options','SAU-012'),(SP,'Options','Options Chain',None),
 (SP,'-','Analyst Ratings','SAU-135'),(SP,'Analyst Ratings','Analyst Ratings Summary',None),
 (SP,'Analyst Ratings Summary','Analyst Ratings Details',None),
 (SP,'Analyst Ratings','Bulls vs Bears','SAU-066'),(SP,'Analyst Ratings','Technical Analysis','SAU-040'),
 (SP,'Analyst Ratings','Whale/Institutional Trades','SAU-133'),
 (SP,'-','Financial',None),(SP,'Financial','Statements',None),(SP,'Financial','Earnings',None),
 (SP,'Financial','Actions',None),
 (SP,'-','Profile','SAU-122'),(SP,'Profile','Major Shareholder','SAU-045'),
 (SP,'-','Holding',None),(SP,'-','Traders',None),(SP,'Traders','Insider Traders','SAU-069'),
 (SP,'Traders','Government Traders','SAU-070'),
 (SP,'-','Share Stock',None),(SP,'-','Add to Watchlist',None),
 (SP,'-','Stock Performance Chart',None),(SP,'Stock Performance Chart','Historical Data',None),
 (SP,'-','Compare Stock','SAU-041'),(SP,'-','AI Assistant','SAU-121'),
 (SP,'-','Stock Key Fact','XJ-058'),(SP,'-','Stock Dividend','SAU-130'),
 (SP,'-','Unusual Activity','SAU-134'),
]
BACKLOG = {'XJ-015':'related-to:Stock Details','SAU-139':'related-to:Stock Performance Chart',
           'XJ-012':'related-to:us-campaign'}
INITIATIVE = {'SAU-121':'initiative:ai-guided-investing'}

wb = openpyxl.load_workbook(X); ws = wb['Features']
L1 = 21
ws.cell(2,28).value='display_order'
I = {str(ws.cell(r,1).value).strip(): r for r in range(4, ws.max_row+1) if ws.cell(r,1).value}
def tags(f): return [t.strip() for t in str(ws.cell(I[f],9).value or '').split(',') if t.strip()]
def settags(f,t): ws.cell(I[f],9).value = ', '.join(t)

log=[]; nxt=max(int(k.split('-')[1]) for k in I if k.startswith('SAU-'))+1
row=max(I.values())+1
resolved={}
for scr,par,name,fid in C:
    if fid:
        r=I[fid]; old=str(ws.cell(r,2).value)
        t=[x for x in tags(fid) if not x.startswith('sub-of:')]
        if old!=name:
            ws.cell(r,2).value=name
            note=f'renamed-from:{old}, v1.4'
            if note not in t: t.append(note)
            log.append(f'  RENAME  {fid:8} {old[:34]!r} -> {name!r}')
        oa=ws.cell(r,10).value
        if oa!='Market': ws.cell(r,10).value='Market'; log.append(f'  AREA    {fid:8} {oa} -> Market')
        if ws.cell(r,11).value!=scr: log.append(f'  SCREEN  {fid:8} {ws.cell(r,11).value} -> {scr}')
        ws.cell(r,11).value=scr
        if fid in INITIATIVE and INITIATIVE[fid] not in t: t.append(INITIATIVE[fid])
        settags(fid,t)
    else:
        fid=f'SAU-{nxt}'; nxt+=1
        ws.cell(row,1).value=fid; ws.cell(row,2).value=name; ws.cell(row,3).value='Saudi Market'
        ws.cell(row,4).value='Planned'; ws.cell(row,9).value='market-ia-v1.4'
        ws.cell(row,10).value='Market'; ws.cell(row,11).value=scr; ws.cell(row,19).value=2
        I[fid]=row; row+=1
        log.append(f'  CREATE  {fid:8} {name!r} [{scr}] Planned')
    resolved[(scr,name)]=fid

haskids={}
for scr,par,name,_ in C:
    if par!='-': haskids[(scr,par)]=True
for scr,par,name,_ in C:
    fid=resolved[(scr,name)]; r=I[fid]
    t=[x for x in tags(fid) if not x.startswith('sub-of:')]
    lv=['ARC Platform','Saudi Market','Market',scr,'','','']
    if par=='-':
        if (scr,name) in haskids: lv[4]=name
        else: lv[5]=name
    else:
        gp=next((p for s,p,n,_ in C if s==scr and n==par), '-')
        if gp=='-':
            lv[4]=par; lv[5]=name
        else:
            lv[4]=gp; lv[5]=par; lv[6]=name
        t.append(f'sub-of:{resolved[(scr,par)]}')
    settags(fid,t)
    for i in range(7): ws.cell(r,L1+i).value=lv[i] or None
    # display_order: contract sequence, depth-first. C is already in outline order,
    # so the index is the order. Every future contract write gets this for free.
    ws.cell(r,28).value=[i for i,(s,_,n,_) in enumerate(C,1) if s==scr and n==name][0]

for fid,tag in BACKLOG.items():
    r=I[fid]; ws.cell(r,11).value=None
    t=[x for x in tags(fid) if not x.startswith('sub-of:')]
    if tag not in t: t.append(tag)
    settags(fid,t)
    for i in range(7): ws.cell(r,L1+i).value=None
    ws.cell(r,28).value=None
    log.append(f'  BACKLOG {fid:8} screen cleared, tagged {tag}')

total=sum(1 for r in range(4,ws.max_row+1) if ws.cell(r,1).value)
ws.cell(3,1).value=(f"Last updated: {date.today().strftime('%d %b %Y')}  |  Schema v1.6  |  {total} features"
                    f"  |  Baseline v1.4 (Portfolio + Market IA)")
wb.save(X)
for l in log: print(l)
print(f'\nsaved — {total} features  ({len(log)} operations)')
