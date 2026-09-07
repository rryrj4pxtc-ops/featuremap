#!/usr/bin/env python3
"""
split_nonarb_areas.py
Replaces Non-ARB area with Local and Global for Onboarding features.
- Local: Non-ARB Saudi entities (corporate, local bank accounts)
- Global: GCC, foreign, international customers
"""
import os, shutil
from datetime import date
from openpyxl import load_workbook

LOCAL = {
    "ONB-019",  # Non-ARB Bank Account Entry
    "ONB-028",  # Corporate — Other Local Entities
    "ONB-029",  # Corporate — Approval Flow
    "ONB-030",  # Corporate — Under Review
    "ONB-031",  # Corporate — Additional Requirements
    "ONB-032",  # Corporate — Rejected
    "ONB-033",  # Corporate — Local Company Onboarding
    "ONB-034",  # Corporate — Profile & Settings
    "ONB-035",  # Corporate — Requirements Upload
    "ONB-036",  # Corporate — Terms & Conditions
    "ONB-037",  # Corporate — Status Tracking
    "ONB-038",  # Corporate — Pending Actions
}

GLOBAL = {
    "ONB-005",  # GCC Customer Onboarding
    "ONB-006",  # Family Onboarding (Non-Saudi)
    "ONB-008",  # GCC Customer ID Update
    "ONB-009",  # Foreign Customer ID Verification
    "ONB-013",  # IBKR Intl Brokerage Onboarding
    "ONB-021",  # Family Members Overview
}

here = os.path.dirname(os.path.abspath(__file__))
dashboard_dir = os.path.join(here, "..")
xlsx_path = os.path.join(dashboard_dir, "features-master.xlsx")
backup_path = xlsx_path + f".bak-{date.today().isoformat()}-split-nonarb"

shutil.copy2(xlsx_path, backup_path)
print(f"Backup: {backup_path}")

wb = load_workbook(xlsx_path)
ws = wb["Features"]

updated = 0
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if not fid:
        continue
    fid = str(fid).strip()
    if fid in LOCAL:
        ws.cell(row=r, column=10, value="Local")
        updated += 1
    elif fid in GLOBAL:
        ws.cell(row=r, column=10, value="Global")
        updated += 1

wb.save(xlsx_path)
wb.close()
print(f"Updated {updated} features: Local={len(LOCAL)}, Global={len(GLOBAL)}")
print(f"Saved: {xlsx_path}")
