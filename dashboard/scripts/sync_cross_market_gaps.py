#!/usr/bin/env python3
"""
sync_cross_market_gaps.py
Full cross-market gap sync:
1. Delete XJ-003 (Guest Mode duplicate — keep ONB-039 only)
2. Add Saudi capabilities missing from US Trading as US Gap features
3. Add US capabilities missing from Saudi Market as Saudi Gap features

Excludes market-specific infrastructure (Nomu, Tadawul, Rights, Murabaha, Sukuk)
and features that already have equivalents in the other journey.
"""
import os, shutil
from datetime import date
from openpyxl import load_workbook

here = os.path.dirname(os.path.abspath(__file__))
dashboard_dir = os.path.join(here, "..")
xlsx_path = os.path.join(dashboard_dir, "features-master.xlsx")
backup_path = xlsx_path + f".bak-{date.today().isoformat()}-sync"

shutil.copy2(xlsx_path, backup_path)
print(f"Backup: {backup_path}")

wb = load_workbook(xlsx_path)
ws = wb["Features"]

# ── Step 1: Delete XJ-003 (Guest Mode — duplicate of ONB-039) ──
deleted_row = None
for r in range(4, ws.max_row + 1):
    fid = ws.cell(row=r, column=1).value
    if fid and str(fid).strip() == "XJ-003":
        deleted_row = r
        break

if deleted_row:
    ws.delete_rows(deleted_row)
    print(f"Deleted XJ-003 (Guest Mode) from row {deleted_row} — ONB-039 is the canonical version")
else:
    print("XJ-003 not found — skipping delete")

# ── Re-read features after delete ──
existing = {}
for row in ws.iter_rows(min_row=4, values_only=True):
    if not any(row[:4]):
        continue
    fid = str(row[0]).strip() if row[0] else None
    if fid:
        existing[fid] = {
            "name": str(row[1]).strip() if row[1] else "",
            "journey": str(row[2]).strip() if row[2] else "",
            "status": str(row[3]).strip() if row[3] else "",
        }

# Get max IDs
us_ids = [int(f.split("-")[1]) for f in existing if f.startswith("US-")]
sau_ids = [int(f.split("-")[1]) for f in existing if f.startswith("SAU-")]
next_us = max(us_ids) + 1   # 040
next_sau = max(sau_ids) + 1  # 108

# ── Step 2: New US Gap features from Saudi capabilities ──
# These are Saudi Live/Planned/Gap features that have NO equivalent in US Trading
# Excluding: Nomu, Tadawul-specific, Rights, Murabaha, Sukuk, and already-mapped features

new_us_gaps = [
    # Trading capabilities
    ("Market-on-Close (MOC) (US)", "Market-on-close order type for US market — execute at closing auction price."),
    ("Short Selling (US)", "Short selling capability for US market positions."),
    ("Attached Order (US)", "Attached order type linking a child order to a parent order for US trades."),
    ("Slicing Order (US)", "Slicing order type to split large US orders into smaller tranches for reduced market impact."),
    ("Chain Order (US)", "Chain order type executing a sequence of US orders one after another."),
    ("Trailing Order (US)", "Trailing stop order that adjusts dynamically with US stock price movement."),
    ("Iceberg Order (US)", "Iceberg order type showing only a portion of total order size for US markets."),
    ("Buying Power Swap (US)", "Swap buying power between US portfolios or accounts."),
    ("Instant Settlement (US)", "Accelerated settlement (T+0 or T+1) for US market trades."),
    # Portfolio
    ("Portfolio P&L View (US)", "Profit and loss overview for US market portfolio holdings."),
    ("Update Avg Cost Price (US)", "Manually adjust average cost price for US holdings to reflect transfers or corporate actions."),
    ("Portfolio Holdings (US)", "Detailed view of all US market holdings with current value and quantity."),
    ("Holdings Letter / PDF (US)", "Generate and download a PDF letter of US portfolio holdings."),
    ("Performance Chart (US)", "Portfolio performance chart tracking US holdings over time."),
    ("Peer Portfolio Comparison (US)", "Compare US portfolio performance against peer investor benchmarks."),
    ("Portfolio Insights (US)", "AI or rules-based insights on US portfolio composition and risk."),
    ("Portfolio Health Score (US)", "Composite health score for US portfolio based on diversification, risk, and performance."),
    ("Allocation Tracker (US)", "Track US portfolio allocation vs target allocation by sector and asset class."),
    ("Goal Tracker (US)", "Track investment goals and progress for US market portfolio."),
    ("Holdings Filter Panel (US)", "Advanced filtering for US holdings by sector, performance, and holding period."),
    ("Liquidate Holdings (US)", "One-action liquidation of selected US market holdings."),
    ("Analysis Dashboard (US)", "Consolidated analysis dashboard for US portfolio metrics and trends."),
    ("Downloadable Reports (US)", "Download US portfolio reports in PDF or Excel format."),
    ("Visual Asset Allocation (US)", "Visual chart showing US portfolio asset allocation breakdown."),
    ("Sector / Index Filters (US)", "Filter US stocks by sector and index membership."),
    # Market intelligence
    ("Sector & Exchange Volatility (US)", "Volatility indicators by sector and exchange for US markets."),
    ("News Screen (US)", "Dedicated news screen for US market news and stock-specific headlines."),
    ("Analyst Ratings (US)", "Consensus analyst ratings and price targets for US stocks."),
    ("Money Flow Tracking (US)", "Track institutional and retail money flow in and out of US stocks."),
    # Watchlist
    ("Watchlist (US)", "Dedicated US market watchlist with real-time prices and daily change."),
    ("Watchlist Story View (US)", "Instagram-style story view for US watchlist stocks with key metrics."),
    ("Minichart Component (US)", "Inline mini price charts within US stock lists and watchlists."),
    ("StoryTeller (US)", "Narrative-driven stock summaries for US watchlist and portfolio holdings."),
    # Social
    ("Social / Copy Trading (US)", "Social trading and copy trading for US market positions."),
    # From Saudi Gaps (should also be US Gaps)
    ("P&L History — Sold Stocks (US)", "Historical profit and loss for previously sold US positions."),
    ("Batch Close Positions (US)", "Close multiple US positions simultaneously in one action."),
    ("Auto Dividend Reinvest (US)", "Automatically reinvest US stock dividends into the same holding."),
    ("Chart Trading (US)", "Execute US trades directly from the price chart interface."),
    ("Chart Touch-and-Hold (US)", "Touch-and-hold gesture on US charts to reveal detailed price data at any point."),
    ("Chart Axis to Market Close (US)", "Extend US chart x-axis to market close to visualize remaining trading day."),
    ("Standard S&D Layout (US)", "Standard supply and demand layout for US stock order book visualization."),
    ("Auto-Watchlist on Buy (US)", "Automatically add US stocks to watchlist when purchased."),
]

# ── Step 3: New Saudi Gap features from US capabilities ──
# US Live/Gap features that have NO equivalent in Saudi Market
# Excluding: Arabic names/news (already Arabic), Bonds (= Sukuk), AI (= Cross-Journey)

new_sau_gaps = [
    ("Pre-Auction Trading", "Trading during Tadawul opening and closing auction periods."),
    ("Fractional Shares", "Buy fractional shares of Saudi-listed stocks for smaller investment amounts."),
    ("ETF Exposure View", "Detailed ETF profile page showing holdings, performance, dividends, and sector exposure for Saudi ETFs."),
    ("Whale / Institutional Tracking", "Track large institutional and whale trades on Tadawul."),
    ("Unusual Activity Alerts", "Alerts for unusual trading volume or price activity on Saudi stocks."),
    ("Futures Trading", "Trading Saudi market futures and derivatives contracts."),
    ("Fast Order (No Confirm)", "Quick order execution without confirmation dialog for experienced traders."),
]

# ── Find last data row ──
last_row = 4
for r in range(ws.max_row, 3, -1):
    if any(ws.cell(row=r, column=c).value for c in range(1, 5)):
        last_row = r
        break

next_row = last_row + 1

# ── Write US Gap features ──
print(f"\nAdding {len(new_us_gaps)} US Trading Gap features (US-{next_us:03d} to US-{next_us + len(new_us_gaps) - 1:03d})")
for i, (name, definition) in enumerate(new_us_gaps):
    fid = f"US-{next_us + i:03d}"
    row_idx = next_row + i
    ws.cell(row=row_idx, column=1, value=fid)
    ws.cell(row=row_idx, column=2, value=name)
    ws.cell(row=row_idx, column=3, value="US Trading")
    ws.cell(row=row_idx, column=4, value="Gap")
    ws.cell(row=row_idx, column=5, value=definition)
    print(f"  + {fid}: {name}")

next_row += len(new_us_gaps)

# ── Write Saudi Gap features ──
print(f"\nAdding {len(new_sau_gaps)} Saudi Market Gap features (SAU-{next_sau:03d} to SAU-{next_sau + len(new_sau_gaps) - 1:03d})")
for i, (name, definition) in enumerate(new_sau_gaps):
    fid = f"SAU-{next_sau + i:03d}"
    row_idx = next_row + i
    ws.cell(row=row_idx, column=1, value=fid)
    ws.cell(row=row_idx, column=2, value=name)
    ws.cell(row=row_idx, column=3, value="Saudi Market")
    ws.cell(row=row_idx, column=4, value="Gap")
    ws.cell(row=row_idx, column=5, value=definition)
    print(f"  + {fid}: {name}")

wb.save(xlsx_path)
wb.close()

print(f"\nSaved: {xlsx_path}")
print(f"Total changes: 1 deleted + {len(new_us_gaps)} US gaps + {len(new_sau_gaps)} Saudi gaps = {1 + len(new_us_gaps) + len(new_sau_gaps)}")
