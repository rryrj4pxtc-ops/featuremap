#!/usr/bin/env python3
"""Detect an unguarded overwrite of features-master.xlsx.

Compares the master against data/master-integrity.json, the fingerprint recorded at the
last guardrail run. Exit 0 = clean, 1 = drift (master changed outside the guardrail).

Run it before starting any edit, and after anything that may have written the file
(a Numbers export, a sync client, a manual save).
"""
import hashlib, json, os, sys
from pathlib import Path

D = Path(__file__).resolve().parents[1]
X = D / 'features-master.xlsx'
MAN = D / 'data' / 'master-integrity.json'

STATUS = D / 'data' / 'status.json'


def write_status(master_md5, verdict):
    """Small file the dashboard header reads: what the master is right now, and when
    we last looked. The header compares master_md5 here against source_md5 in
    features_derived.json to decide whether the page is showing stale data."""
    import datetime
    try:
        STATUS.write_text(json.dumps({
            'master_md5': master_md5,
            'checked_at': datetime.datetime.now().astimezone().isoformat(timespec='seconds'),
            'integrity': verdict,
        }, indent=2))
    except Exception:
        pass


if not MAN.exists():
    print('no manifest — run record-integrity to establish a baseline'); sys.exit(1)
m = json.loads(MAN.read_text())

md5 = hashlib.md5(X.read_bytes()).hexdigest()
size = os.path.getsize(X)
drift = []
if md5 != m['md5']:   drift.append(f"md5   {m['md5']} -> {md5}")
if size != m['size']: drift.append(f"size  {m['size']} -> {size} bytes")

try:
    import openpyxl
    ws = openpyxl.load_workbook(X)['Features']
    rows = sum(1 for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value)
    if rows != m['feature_rows']:
        drift.append(f"rows  {m['feature_rows']} -> {rows}")
    hdr = [ws.cell(2, c).value for c in range(1, 33)]   # A..AF incl. Level8, content_spec, row_type
    if hdr != m['headers']:
        drift.append('headers changed')
    counts = {}
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 3).value == 'Saudi Market' and ws.cell(r, 11).value:
            k = ws.cell(r, 11).value; counts[k] = counts.get(k, 0) + 1
    for k, v in m['screen_counts'].items():
        if counts.get(k, 0) != v:
            drift.append(f"screen '{k}'  {v} -> {counts.get(k, 0)} rows")
except Exception as e:
    drift.append(f'could not read workbook: {e}')

if not drift:
    write_status(md5, 'ok')
    print(f"INTEGRITY OK — master matches the {m['baseline']} fingerprint recorded {m['recorded']}")
    sys.exit(0)

write_status(md5, 'drift')

print(f"INTEGRITY DRIFT — master changed outside the guardrail (baseline {m['baseline']}):")
for d in drift: print('   ', d)
print(f"\nRestore with:\n    cp '{m['snapshot']}' features-master.xlsx")
print("Then alert Ahmed before doing anything else. Do NOT regenerate JSON from a drifted master.")
sys.exit(1)
